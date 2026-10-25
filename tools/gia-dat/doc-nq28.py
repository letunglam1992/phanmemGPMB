"""
Đọc phụ biểu NQ 28/2026/NQ-HĐND (sửa đổi, bổ sung bảng giá đất NQ 152/2025/NQ-HĐND) từ tệp Excel do đơn vị cung cấp
→ policy/nguon/nq28-2026-sua-bang-gia-dat.json. Không đưa tệp Excel vào kho; chạy lại khi có bản ký ban hành:

    python3 -I tools/gia-dat/doc-nq28.py <phu-bieu.xlsx> <ten-tep-goc>

Đơn vị: nghìn đồng/m². Phụ lục I (Bảng 05 đất ở), II (Bảng 06 TMDV), III (Bảng 07 SKC) — bổ sung; IV — điều chỉnh tên;
V — bãi bỏ (Điều 1 NQ 28). Trang "Bổ sung NTD" (nghĩa trang) không có trong Điều 1 → ghi vào "chua_ap_dung".
"""
import json, re, sys, os, difflib
import openpyxl

xlsx, ten_goc = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(xlsx, data_only=True)
GOC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "policy", "nguon")
nq152 = json.load(open(os.path.join(GOC, "nq152-2025-bang-gia-dat.json"), encoding="utf-8"))
DM_XA = nq152["danh_muc_xa"]

def sach(s):
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip() if s is not None else ""

def ten_xa(s):
    t = sach(s).lower()
    for x in DM_XA:
        if x.lower() in (t, f"xã {t}", f"phường {t}"):
            return x
    return None

def so(v):
    if v is None or v == "" or isinstance(v, str):
        return None
    return int(v) if float(v).is_integer() else float(v)

def doc_bo_sung(ws, bang, loai):
    out, xa, nhom, n = [], None, "", 0
    for r in ws.iter_rows(min_row=6, values_only=True):
        stt, ten = sach(r[0]), sach(r[1])
        if not ten:
            continue
        vt = [so(v) for v in r[2:7]]
        ghi = sach(r[7]) if len(r) > 7 else ""
        x = ten_xa(ten)
        if x and all(v is None for v in vt):
            xa, nhom = x, ""
            continue
        if all(v is None for v in vt):
            nhom = ten
            continue
        if xa is None:
            raise SystemExit(f"{ws.title}: dòng có giá trước khi có xã: {ten}")
        n += 1
        out.append({"bang": bang, "loai_dat": loai, "xa": xa, "stt": f"NQ28-{n}", "nhom": nhom or "Bổ sung theo NQ 28/2026/NQ-HĐND",
                    "tuyen": ten, "vt": vt, "so_vi_tri": sum(1 for v in vt if v is not None), "canh_bao": [],
                    **({"vi_tri_bo_sung": ghi} if ghi else {}), **({"stt_phu_bieu": stt} if stt and stt != "-" else {})})
    return out

def doc_danh_sach(ws, co_ten_moi):
    out, xa = [], None
    for r in ws.iter_rows(min_row=4, values_only=True):
        stt, ten = sach(r[0]), sach(r[1])
        if not ten:
            continue
        x = ten_xa(ten)
        if x and re.fullmatch(r"5\.\d+", stt):
            xa = x
            continue
        if not stt or not re.fullmatch(r"\d+(\.\d+)?", stt) or xa is None:
            continue
        if co_ten_moi:
            moi, ghi = sach(r[2]), sach(r[3])
            if "." not in stt or not moi:  # dòng nhóm (vd. "1 Quốc lộ 43…", "1.1 Đường Tây Bắc") — không phải tuyến đổi tên
                continue
            out.append({"xa": xa, "stt": stt, "ten_cu": ten, "ten_moi": moi, "ghi_chu": ghi})
        else:
            ghi = sach(r[2])
            out.append({"xa": xa, "stt": stt, "ten": ten, "ghi_chu": ghi})
    return out

BANG = {"05": nq152["dat_o"], "06": nq152["dat_tmdv"], "07": nq152["dat_skc"]}
def gan(a, b):
    return difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio()

def khop(xa, stt, ten):
    """Tìm dòng NQ 152 (từng bảng 05, 06, 07): cùng xã, cùng STT và tên gần giống; không có thì cùng xã, tên trùng (≥ 0,95)."""
    kq = {}
    for b, ds in BANG.items():
        cung = [r for r in ds if r["xa"] == xa]
        m = [r for r in cung if r["stt"] == stt and gan(r["tuyen"], ten) > 0.6]
        if len(m) != 1:
            m = [r for r in cung if gan(r["tuyen"], ten) >= 0.95]
        if len(m) == 1:
            kq[b] = {"stt": m[0]["stt"], "tuyen": m[0]["tuyen"]}
    return kq

def gan_khop(ds, khoa_ten, ten_loai):
    ok, chua = [], []
    for d in ds:
        k = khop(d["xa"], d["stt"], d[khoa_ten])
        if k:
            d["khop"] = k
            ok.append(d)
        elif ten_loai == "bai_bo" and "." not in d["stt"]:
            continue  # dòng tên nhóm trong phụ lục (các tuyến con bãi bỏ ghi riêng)
        else:
            chua.append({**d, "loai": ten_loai})
    return ok, chua

ten_trang = {ws.title: ws for ws in wb.worksheets}
def trang(dau):
    for k, v in ten_trang.items():
        if k.startswith(dau):
            return v
    raise SystemExit(f"Thiếu trang {dau}")

kq = {
    "van_ban": "NQ 28/2026/NQ-HĐND ngày 19/10/2026",
    "sua_doi": "NQ 152/2025/NQ-HĐND ngày 29/12/2025",
    "hieu_luc_tu": "2026-11-01",
    "don_vi": "nghìn đồng/m²",
    "nguon": ten_goc,
    "luu_y_nguon": "Theo tệp đơn vị cung cấp (tên tệp ghi \"dự thảo\"; phụ biểu \"kèm DTNQ update 2/10/2026\") — đối chiếu với bản Nghị quyết đã ký, công bố trước khi áp dụng.",
    "bo_sung": {
        "dat_o": doc_bo_sung(trang("PL.1"), "05", "ODT/ONT"),
        "dat_tmdv": doc_bo_sung(trang("PL.2"), "06", "TMD"),
        "dat_skc": doc_bo_sung(trang("PL.3"), "07", "SKC"),
    },
    "doi_ten": [],
    "bai_bo": [],
    "quy_dinh": [
        {"vi_tri": "khoản 8 Điều 7 (bổ sung bởi khoản 6 Điều 1 NQ 28/2026)", "noi_dung": "Đối với diện tích đất thương mại, dịch vụ của hộ gia đình, cá nhân đang sử dụng ổn định lâu dài có nguồn gốc từ đất ở (được công nhận, giao đất ban đầu là đất ở hoặc được phép chuyển mục đích sử dụng đất từ đất ở) thì giá đất được xác định bằng 90% đơn giá đất ở cùng vị trí, cùng tuyến đường quy định tại Bảng số 05 ban hành kèm theo Quy định này."},
        {"vi_tri": "khoản 3 Điều 4 (bãi bỏ bởi Điều 2 NQ 28/2026)", "noi_dung": "Bãi bỏ khoản 3 Điều 4 Quy định ban hành kèm theo Nghị quyết số 152/2025/NQ-HĐND (giảm giá do chênh lệch độ cao)."},
        {"vi_tri": "khoản 1 Điều 3 NQ 28/2026 (chuyển tiếp)", "noi_dung": "Đối với trường hợp đã có quyết định phê duyệt phương án bồi thường, hỗ trợ và tái định cư trước ngày Nghị quyết này có hiệu lực thi hành thì tiếp tục thực hiện theo phương án đã được phê duyệt, trừ trường hợp phương án bồi thường, hỗ trợ, tái định cư phải điều chỉnh, bổ sung hoặc hủy bỏ theo quy định của pháp luật."},
    ],
    "chua_ap_dung": [
        {"trang": t, "ly_do": "Phụ biểu có trang này nhưng Điều 1 NQ 28/2026 không có nội dung tương ứng (bổ sung giá đất nghĩa trang, nhà tang lễ, cơ sở hỏa táng) — chưa đưa vào bảng giá; chờ đối chiếu bản ký."}
        for t in ten_trang if "NTD" in t
    ],
}
kq["doi_ten"], c1 = gan_khop(doc_danh_sach(trang("PL.4_Điều chỉnh"), True), "ten_cu", "doi_ten")
kq["bai_bo"], c2 = gan_khop(doc_danh_sach(trang("PL.5"), False), "ten", "bai_bo")
for d in kq["doi_ten"]:
    if d["ten_moi"].lower().startswith("sửa đổi stt"):
        d["chi_doi_stt"] = True
kq["chua_khop"] = c1 + c2
ra = os.path.join(GOC, "nq28-2026-sua-bang-gia-dat.json")
json.dump(kq, open(ra, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(ra, {k: len(v) for k, v in kq["bo_sung"].items()}, "đổi tên", len(kq["doi_ten"]), "bãi bỏ", len(kq["bai_bo"]), "chưa khớp", [(x["loai"], x["xa"], x["stt"]) for x in kq["chua_khop"]])
for d in kq["doi_ten"] + kq["bai_bo"]:
    if len(d["khop"]) < 3:
        print("chỉ khớp một phần bảng:", d["xa"], d["stt"], list(d["khop"]))
