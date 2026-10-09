"""
Chuyển 3 mẫu văn bản riêng của địa phương (người dùng cung cấp, dạng .doc → .docx) thành mẫu tự điền:
  - rieng-to-trinh-thu-hoi.docx   (Tờ trình đề nghị thu hồi đất — Phòng Kinh tế)
  - rieng-bao-cao-tham-dinh.docx  (Báo cáo thẩm định thu hồi đất — Phòng Kinh tế)
  - rieng-qd-thu-hoi.docx         (Quyết định thu hồi đất cả đợt, kèm danh sách — Chủ tịch UBND xã)

Giữ nguyên định dạng gốc (bảng tiêu đề, đường kẻ, cỡ chữ, bảng danh sách 17 cột); thay mọi số liệu,
tên người, tên dự án bằng trường {…}. Tệp gốc có thông tin cá nhân → KHÔNG đưa vào kho mã.

Chạy: python3 tools/mau-van-ban/mau-rieng.py <thư mục chứa 3 tệp .docx đã gộp run>

Thư mục nguồn phải có thêm tệp thay-the.json (KHÔNG đưa vào kho mã — chứa tên người ký, số văn bản
của hồ sơ gốc), dạng:
  {"xa_ubnd": "UBND XÃ …", "xa": "XÃ …", "phong": "PHÒNG …",
   "so_tt": "Số: …/TTr-…", "so_bc": "Số: …/BC-…",
   "ky_tt": ["TRƯỞNG PHÒNG", "<họ tên>"], "ky_bc": ["KT. TRƯỞNG PHÒNG", "<họ tên>", "PHÓ TRƯỞNG PHÒNG"],
   "can_bo_td": "<họ tên>", "ky_qd": ["CHỦ TỊCH", "<họ tên>"],
   "cam": ["<chuỗi không được còn trong mẫu sau khi dựng: tên người, số văn bản, email…>"]}
"""
import copy
import json
import warnings

warnings.filterwarnings("ignore", message="Duplicate name")
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tien_ich_mau import dat_doan, dong_goi_lai, khoi_phap_ly, lam_sach, lap, noi_nhan_va_ky, t_elems, than, thay_chu, xoa  # noqa: E402

NGUON = Path(sys.argv[1])
RA = Path(__file__).resolve().parents[2] / "apps/desktop/public/mau-van-ban"
TT = json.loads((NGUON / "thay-the.json").read_text(encoding="utf8"))
PV = "{#co_pham_vi} ({pham_vi_dot}){/co_pham_vi}"


def khoi_dien_tich(ds, phap_ly=False):
    """ds: 10 đoạn của khối diện tích (1. Tổng … đến + Đất đồi…). Dựng lại theo trường; phap_ly: thêm khối theo
    tình trạng pháp lý nguồn gốc đất (tờ trình, báo cáo thẩm định — không đưa vào quyết định)."""
    p_tong, p_duoc, p_ho, p_ho_l1, p_ho_l2, p_tc, p_khong, p_k1, p_k2, p_k3 = ds
    dat_doan(p_tong, [("1. Tổng diện tích đất thu hồi: {tong_dt_thu_hoi} m2, trong đó:", False)])
    dat_doan(p_duoc, [("* Diện tích đất được bồi thường, hỗ trợ: {dt_duoc_bt} m2.", False)])
    dat_doan(p_ho, [("- Đất của hộ gia đình, cá nhân: {dt_duoc_bt_ho} m2{gom_ho}", False)])
    lap(p_ho_l1, "dt_duoc_bt_ho_loai", [("+ {ten}: {dien_tich} m2.", False)])
    dat_doan(p_tc, [("- Đất của tổ chức: {dt_duoc_bt_tc} m2{gom_tc}", False)])
    p_tc_l = copy.deepcopy(p_ho_l1)
    p_tc.addnext(p_tc_l)
    dat_doan(p_tc_l, [("+ {ten}: {dien_tich} m2.", False)])
    lap(p_tc_l, "dt_duoc_bt_tc_loai", [("+ {ten}: {dien_tich} m2.", False)])
    dat_doan(p_khong, [("* Tổng diện tích đất không được bồi thường, hỗ trợ: {dt_khong_bt} m2{gom_khong}", False)])
    lap(p_k1, "dt_khong_bt_loai", [("+ {ten}: {dien_tich} m2.", False)])
    xoa(p_ho_l2, p_k2, p_k3)
    if phap_ly:
        khoi_phap_ly(p_k1.getnext())


def bang_danh_sach(tbl):
    """Bảng 17 cột: giữ 3 hàng tiêu đề, hàng tổng (bản), 1 hàng dữ liệu lặp; bỏ các hàng còn lại."""
    trs = tbl.findall(qn("w:tr"))
    tong, mau = trs[3], trs[4]
    xoa(*trs[5:])
    TRUONG = ["stt", "ho_ten", "so_to", "so_thua", "dt_thu_hoi", "loai_dat", "gcn_seri", "gcn_to", "gcn_thua", "gcn_dien_tich", "gcn_loai_dat", "dt_co_gcn", "loai_dat_co_gcn", "dt_khong_gcn", "loai_dat_khong_gcn", "nguon_goc", "ghi_chu"]
    tcs = mau.findall(qn("w:tc"))
    assert len(tcs) == 17, len(tcs)
    for i, (tc, f) in enumerate(zip(tcs, TRUONG)):
        p = tc.find(qn("w:p"))
        for p2 in tc.findall(qn("w:p"))[1:]:
            xoa(p2)
        s = ("{#ds_thua_thu_hoi}" if i == 0 else "") + "{" + f + "}" + ("{/ds_thua_thu_hoi}" if i == 16 else "")
        dat_doan(p, [(s, False)])
    tong_tr = {1: "{ban_khu_dan_cu_bang}", 4: "{tong_dt_thu_hoi}", 11: "{tong_dt_co_gcn}", 13: "{tong_dt_khong_gcn}"}
    for i, tc in enumerate(tong.findall(qn("w:tc"))):
        p = tc.find(qn("w:p"))
        for p2 in tc.findall(qn("w:p"))[1:]:
            xoa(p2)
        dat_doan(p, [(tong_tr.get(i, ""), True)])


def to_trinh(src):
    d = Document(src)
    b = than(d)
    hdr = b[0]
    thay_chu(hdr, TT["xa_ubnd"], "UBND {TEN_XA}")
    thay_chu(hdr, TT["phong"], "{TEN_PHONG}")
    thay_chu(hdr, TT["so_tt"], "Số: {so}/TTr-{ky_hieu_phong}")
    ngay = next(t for t in t_elems(hdr) if "năm 2026" in (t.text or ""))
    ngay.text = "{dia_danh}, ngày {ngay} tháng {thang} năm {nam}"
    dat_doan(b[3], [("Đề nghị thu hồi đất để thực hiện dự án: {ten_du_an}", True)])
    xoa(b[4], b[5])
    dat_doan(b[6], [("{#co_pham_vi}({pham_vi_dot}){/co_pham_vi}", True)])
    dat_doan(b[8], [("Kính gửi: Chủ tịch Ủy ban nhân dân {ten_xa}.", False)])
    lap(b[10], "can_cu", [("{.}", False)])
    xoa(*b[11:19])
    dat_doan(b[19], [("Thực hiện Thông báo số {tb_thu_hoi_so} ngày {tb_thu_hoi_ngay} của UBND {ten_xa} về việc thông báo thu hồi đất dự án {ten_du_an}" + PV + ";", False)])
    dat_doan(b[20], [("Theo đề nghị của {ten_don_vi_bt} tại Tờ trình số {tt_don_vi_so} ngày {tt_don_vi_ngay} về việc đề nghị thu hồi đất để thực hiện dự án: {ten_du_an}" + PV + ";", False)])
    dat_doan(b[21], [("{ten_phong} kính đề nghị Chủ tịch UBND {ten_xa} thu hồi đất của {so_doi_tuong_mo_ta} để thực hiện dự án: {ten_du_an}" + PV + ", với các nội dung như sau:", False)])
    khoi_dien_tich(b[22:32], phap_ly=True)
    dat_doan(b[32], [("2. Địa điểm:", True), (" {dia_diem_du_an}.", False)])
    dat_doan(b[33], [("3. Lý do thu hồi đất:", True), (" {ly_do_thu_hoi}.", False)])
    dat_doan(b[35], [("{ten_phong} kính trình Chủ tịch UBND {ten_xa} xem xét, quyết định./.", False)])
    noi_nhan_va_ky(b[36], *TT["ky_tt"])
    dat_doan(b[40], [("Dự án: {ten_du_an}" + PV, True)])
    dat_doan(b[41], [("(Kèm theo Tờ trình số {so}/TTr-{ky_hieu_phong} ngày {ngay_ky_ngan} của {ten_phong} {ten_xa})", False)])
    bang_danh_sach(b[42])
    lam_sach(d)
    d.save(RA / "rieng-to-trinh-thu-hoi.docx")


def bao_cao(src):
    d = Document(src)
    b = than(d)
    hdr = b[0]
    thay_chu(hdr, TT["xa_ubnd"], "UBND {TEN_XA}")
    thay_chu(hdr, TT["phong"], "{TEN_PHONG}")
    thay_chu(hdr, TT["so_bc"], "Số: {so}/BC-{ky_hieu_phong}")
    ngay = next(t for t in t_elems(hdr) if "năm 2026" in (t.text or ""))
    ngay.text = "{dia_danh}, ngày {ngay} tháng {thang} năm {nam}"
    dat_doan(b[3], [("Thẩm định thu hồi đất dự án: {ten_du_an}", True)])
    xoa(b[4], b[5])
    dat_doan(b[6], [("{#co_pham_vi}({pham_vi_dot}){/co_pham_vi}", True)])
    dat_doan(b[8], [("Kính gửi: Chủ tịch Ủy ban nhân dân {ten_xa}.", False)])
    lap(b[11], "can_cu_gach", [("- {.}", False)])
    xoa(*b[12:23])
    dat_doan(b[25], [("- Tờ trình số {tt_don_vi_so} ngày {tt_don_vi_ngay} của {ten_don_vi_bt} về việc đề nghị thu hồi đất để thực hiện dự án: {ten_du_an}" + PV + ".", False)])
    khoi_dien_tich(b[28:38], phap_ly=True)
    dat_doan(b[39], [("Nhất trí theo Tờ trình số {tt_don_vi_so} ngày {tt_don_vi_ngay} của {ten_don_vi_bt}, cụ thể như sau:", False)])
    khoi_dien_tich(b[40:50])
    dat_doan(b[50], [("3. Địa điểm:", True), (" {dia_diem_du_an}.", False)])
    dat_doan(b[51], [("4. Kết luận:", True), (" {ket_luan_tham_dinh}", False)])
    dat_doan(b[52], [("Trên đây là kết quả thẩm định thu hồi đất thực hiện dự án {ten_du_an}" + PV + "./.", False)])
    noi_nhan_va_ky(b[53], TT["ky_bc"][0], TT["ky_bc"][1], bo_dong=TT["ky_bc"][2])
    thay_chu(b[53], TT["can_bo_td"], "{can_bo_tham_dinh}")
    lam_sach(d)
    d.save(RA / "rieng-bao-cao-tham-dinh.docx")


def quyet_dinh(src):
    d = Document(src)
    b = than(d)
    hdr = b[0]
    thay_chu(hdr, TT["xa"], "{TEN_XA}")
    so = next(t for t in t_elems(hdr) if (t.text or "").strip().startswith("Số:"))
    so.text = "Số: {so}/QĐ-UBND"
    ngay = next(t for t in t_elems(hdr) if "năm 2026" in (t.text or ""))
    ngay.text = "{dia_danh}, ngày {ngay} tháng {thang} năm {nam}"
    dat_doan(b[3], [("Thu hồi đất để thực hiện dự án: {ten_du_an}", True)])
    xoa(b[4], b[5])
    dat_doan(b[6], [("{#co_pham_vi}({pham_vi_dot}){/co_pham_vi}", True)])
    dat_doan(b[8], [("CHỦ TỊCH ỦY BAN NHÂN DÂN {TEN_XA}", True)])
    lap(b[10], "can_cu", [("{.}", False)])
    xoa(*b[11:19])
    dat_doan(b[19], [("Theo đề nghị của {chuc_danh_de_nghi} tại Tờ trình số {to_trinh_so} ngày {to_trinh_ngay}.", False)])
    dat_doan(b[21], [("Điều 1.", True), (" Thu hồi đất của {so_doi_tuong_mo_ta} để thực hiện dự án: {ten_du_an}" + PV + ", với các nội dung sau:", False)])
    khoi_dien_tich(b[22:32])
    dat_doan(b[32], [("2. Địa điểm:", True), (" {dia_diem_du_an}.", False)])
    dat_doan(b[33], [("3. Lý do thu hồi đất:", True), (" {ly_do_thu_hoi}.", False)])
    dat_doan(b[36], [("1.", True), (" {ten_don_vi_bt} chủ trì, phối hợp với {ten_phong} {ten_xa} giao Quyết định này đến các chủ sử dụng đất; trường hợp chủ sử dụng đất không nhận Quyết định này hoặc vắng mặt thì phải lập biên bản; niêm yết Quyết định này tại {dia_diem_niem_yet_qd} trong thời gian thực hiện dự án; tổ chức thu hồi giấy chứng nhận quyền sử dụng đất (nếu có) chuyển đến {co_quan_chinh_ly} để chỉnh lý hồ sơ địa chính theo quy định.", False)])
    dat_doan(b[37], [("2.", True), (" {ten_don_vi_bt} có trách nhiệm chủ trì, phối hợp với các cơ quan, đơn vị liên quan thực hiện bàn giao đất cho đơn vị thi công công trình thực hiện dự án.", False)])
    dat_doan(b[38], [("3.", True), (" {co_quan_chinh_ly} chỉnh lý hồ sơ địa chính theo quy định.", False)])
    dat_doan(b[39], [("4.", True), (" {co_quan_dang_tai} có trách nhiệm đăng tải Quyết định này trên Trang thông tin điện tử của xã.", False)])
    dat_doan(b[41], [("1. Quyết định này có hiệu lực từ ngày {ngay_hieu_luc}.", False)])
    noi_nhan_va_ky(b[44], *TT["ky_qd"])
    dat_doan(b[57], [("Dự án: {ten_du_an}", True)])
    xoa(b[58])
    dat_doan(b[59], [("{#co_pham_vi}({pham_vi_dot}){/co_pham_vi}", True)])
    dat_doan(b[60], [("(Kèm theo Quyết định số {so}/QĐ-UBND ngày {ngay_ky_ngan} của Chủ tịch UBND {ten_xa})", False)])
    bang_danh_sach(b[61])
    lam_sach(d)
    d.save(RA / "rieng-qd-thu-hoi.docx")


def tim(mau):
    return next(NGUON.glob(mau))


to_trinh(tim("*To_trinh*.docx"))
bao_cao(tim("*Bao_cao*.docx"))
quyet_dinh(tim("*QD_thu_hoi*.docx"))
for f in ["rieng-to-trinh-thu-hoi.docx", "rieng-bao-cao-tham-dinh.docx", "rieng-qd-thu-hoi.docx"]:
    dong_goi_lai(RA / f)

# Kiểm tra: không còn tên người, số liệu của hồ sơ gốc
CAM = TT["cam"] + ["@"]
for f in ["rieng-to-trinh-thu-hoi.docx", "rieng-bao-cao-tham-dinh.docx", "rieng-qd-thu-hoi.docx"]:
    import zipfile
    z = zipfile.ZipFile(RA / f)
    anh = [n for n in z.namelist() if n.startswith("word/media/")]
    chu = "\n".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    con = [c for c in CAM if re.search(re.escape(c), re.sub(r"<[^>]+>", " ", chu))]
    loi = (f"CÒN: {con} " if con else "") + (f"ẢNH: {anh}" if anh else "")
    print(f, "OK" if not loi else loi)
    if loi:
        raise SystemExit(1)
